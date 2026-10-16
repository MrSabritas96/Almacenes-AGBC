from openpyxl.utils import get_column_letter
import openpyxl,json,re,hashlib,unicodedata,datetime,sys
from pathlib import Path
if len(sys.argv)!=3: raise SystemExit('Uso: python extraer-archivo-agbc.py CARPETA_ACTAS SALIDA_JSON')
root=Path(sys.argv[1])
def norm(v):return ''.join(c for c in unicodedata.normalize('NFD',str(v or '').upper()) if not unicodedata.combining(c)).strip()
months={m:i+1 for i,m in enumerate(['ENERO','FEBRERO','MARZO','ABRIL','MAYO','JUNIO','JULIO','AGOSTO','SEPTIEMBRE','OCTUBRE','NOVIEMBRE','DICIEMBRE'])}
files=[]
for fn in sorted(p.name for p in root.glob('*.xlsx') if not p.name.startswith(('Base ', '~' + '$'))):
 p=root/fn; w=openpyxl.load_workbook(p,read_only=True,data_only=True); docs=[]; skipped=[]
 for ws in w:
  cells={c.coordinate:str(c.value) for row in ws for c in row if c.value is not None}
  titles=[(c.coordinate,c.row,c.column,str(c.value)) for row in ws for c in row if isinstance(c.value,str) and re.search(r'ACTA DE (ASIGN|ACTUAL|DEVOL)',norm(c.value))]
  if not titles:skipped.append(ws.title);continue
  for coord,row,col,title in titles:
   right=min([cc for _,_,cc,_ in titles if cc>col] or [min(ws.max_column+1,col+6)])
   end=min([rr for _,rr,cc,_ in titles if rr>row and cc==col] or [ws.max_row+1])
   def val(rr,cc):return cells.get(f'{get_column_letter(cc)}{rr}','').strip()
   def after(label):
    for rr in range(row+1,min(end,row+13)):
     if norm(val(rr,col)).startswith(label):
      return next((val(k,col) for k in range(rr+1,min(end,rr+4)) if val(k,col)),'')
    return ''
   reference=next((val(rr,col) for rr in range(row+1,min(end,row+5)) if re.search(r'\d+.*?/\s*20\d{2}',norm(val(rr,col)))),'')
   responsible=after('RESPONSABLE');cargo=after('CARGO');narrative=next((val(rr,col) for rr in range(row+1,min(end,row+15)) if norm(val(rr,col)).startswith('EN LA CIUDAD')),'')
   match=re.search(r'(\d{1,2})\s+DE\s+([A-Z]+)\s+(?:DE(?:L)?\s+)?(20\d{2})',norm(narrative));fecha=None
   if match and match[2] in months:
    try:fecha=datetime.date(int(match[3]),months[match[2]],int(match[1])).isoformat()
    except ValueError:pass
   city=re.search(r'EN LA CIUDAD DE (.*?),',narrative,re.I)
   bienes=[]; header=next((rr for rr in range(row+1,min(end,row+20)) if 'DESCRIPCI' in norm(val(rr,col+1))),None)
   if header:
    for rr in range(header+1,end):
     if not re.fullmatch(r'\d+(?:\.0)?',val(rr,col)):continue
     description=val(rr,col+1)
     if not description:continue
     vals=[val(rr,cc) for cc in range(col+2,right)]
     estado=next((v for v in vals if norm(v) in ['BUENO','REGULAR','MALO','MAL ESTADO']),'')
     codes=[v for v in vals[:2] if v and v!=estado]
     bienes.append({'fila':rr,'numero':val(rr,col),'descripcion':description,'codigos':codes,'estado':estado,'observaciones':' | '.join(vals[3:]) if len(vals)>3 else ''})
   title_key=norm(title)
   procedencia=next((v for k,v in [('EX-ECOBOL','EXECOBOL 2024'),('EX ECOBOL','EXECOBOL 2024'),('ECOBOL','EXECOBOL 2024'),('UPAEP','UPAEP'),('UPAED','UPAEP'),('MOPSV','MOPSV'),('PACTA','PACTA'),('UPU','UPU'),('AGBC','AGBC')] if k in title_key),'POR REVISAR')
   docs.append({'procedencia':procedencia,'hoja':ws.title,'bloque':coord,'titulo':title,'tipo':'devolucion' if 'DEVOL' in norm(title) else 'actualizacion' if 'ACTUAL' in norm(title) else 'asignacion','numero':reference,'responsable':responsible,'cargo':cargo,'fecha':fecha,'ciudad':city[1].strip() if city else '', 'texto':narrative,'bienes':bienes,'celdas':cells,'incidencias':(['Procedencia por revisar'] if procedencia=='POR REVISAR' else [])+(['Sin número identificado'] if not reference else [])+(['Sin responsable identificado'] if not responsible else [])+(['Fecha por revisar'] if not fecha else [])+(['Sin bienes identificados'] if not bienes else [])})
 files.append({'nombre':fn,'ruta':str(p),'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'hojas':len(w.sheetnames),'omitidas':skipped,'actas':docs});w.close()
Path(sys.argv[2]).write_text(json.dumps(files,ensure_ascii=False,indent=2),encoding='utf-8')
for f in files: print(json.dumps({'archivo':f['nombre'],'hojas':f['hojas'],'actas':len(f['actas']),'lineasBienes':sum(len(a['bienes']) for a in f['actas']),'omitidas':f['omitidas'],'conIncidencias':sum(bool(a['incidencias']) for a in f['actas'])},ensure_ascii=False))

